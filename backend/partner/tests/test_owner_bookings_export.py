"""The hotel owner downloads the bookings of their own hotels as CSV or Excel, with no guest names, e-mails or phones."""
import csv
import io
from datetime import date, time
from decimal import Decimal

import pytest
from django.test import override_settings
from openpyxl import load_workbook
from rest_framework.test import APIClient

from admin_panel.models import AdminAccessLog
from bookings.tests.r12_helpers import make_booking
from common.dates import business_time_zone
from partner.views import BOOKINGS_EXPORT_HEADER
from promotions.tests.helpers import make_guest, make_hotel, make_owner, make_staff
from properties.models import PropertyTranslation
from users.models import User

pytestmark = pytest.mark.django_db

URL = '/api/v1/partner/bookings/'


def client_of(user):
    api = APIClient()
    api.force_authenticate(user=user)
    return api


@pytest.fixture
def owner():
    return make_owner('mine@example.com')


@pytest.fixture
def hotel(owner):
    return make_hotel(owner, name='My Hotel')


@pytest.fixture
def guest():
    return User.objects.create_user(
        email='guest.private@example.com', password='x', first_name='Private', last_name='Person')


def rows_of(response):
    text = b''.join(response.streaming_content).decode('utf-8-sig')
    return list(csv.reader(io.StringIO(text)))


def book(hotel, guest, check_in, **extra):
    return make_booking(hotel, guest, check_in, **extra)


class TestTheCsvFile:
    def test_one_row_per_booking_newest_check_in_first_with_the_agreed_columns(self, owner, hotel, guest):
        older = book(hotel, guest, date(2026, 11, 3), nights=2, rooms=2, guest_count=3, price=Decimal('100000'))
        newer = book(hotel, guest, date(2026, 12, 1), nights=1, status='cancelled', payment_status='pending')
        response = client_of(owner).get(URL, {'export': 'csv'})
        assert response.status_code == 200
        assert response['Content-Type'].startswith('text/csv')
        assert response['Content-Disposition'] == 'attachment; filename="partner_bookings.csv"'
        rows = rows_of(response)
        assert rows[0] == BOOKINGS_EXPORT_HEADER
        by_reference = {row[0]: row for row in rows[1:]}
        assert [row[0] for row in rows[1:]] == [newer.confirmation_code, older.confirmation_code]
        line = dict(zip(BOOKINGS_EXPORT_HEADER, by_reference[older.confirmation_code]))
        assert (line['hotel'], line['status'], line['check_in'], line['check_out'], line['nights']) == (
            'My Hotel', 'confirmed', '2026-11-03', '2026-11-05', '2')
        assert (line['rooms'], line['guests'], line['total'], line['currency']) == ('2', '3', '400000.00', 'UZS')
        assert line['booked_on'] == older.created_at.astimezone(business_time_zone()).date().isoformat()

    def test_the_file_carries_no_guest_names_emails_or_phones(self, owner, hotel, guest):
        booking = book(hotel, guest, date(2026, 11, 3))
        booking.guest_phone = '+998901234567'
        booking.special_requests = 'Quiet room please'
        booking.save()
        text = b''.join(client_of(owner).get(URL, {'export': 'csv'}).streaming_content).decode('utf-8-sig')
        for private in ('Test Guest', 'guest.private@example.com', 'Private', 'Person', '+998901234567', 'Quiet room'):
            assert private not in text

    def test_the_arrival_time_and_when_the_guest_arrived(self, owner, hotel, guest):
        booking = book(hotel, guest, date(2026, 11, 3))
        booking.expected_arrival_time = time(14, 30)
        booking.save()
        line = dict(zip(BOOKINGS_EXPORT_HEADER, rows_of(client_of(owner).get(URL, {'export': 'csv'}))[1]))
        assert (line['arrival_time'], line['arrived']) == ('14:30', 'no')

    def test_only_the_bookings_of_the_owners_own_hotels(self, owner, hotel, guest):
        mine = book(hotel, guest, date(2026, 11, 3))
        stranger_hotel = make_hotel(make_owner('other@example.com'), name='Other Hotel')
        book(stranger_hotel, guest, date(2026, 11, 3))
        rows = rows_of(client_of(owner).get(URL, {'export': 'csv'}))
        assert [row[0] for row in rows[1:]] == [mine.confirmation_code]

    def test_the_filters_of_the_list_apply_to_the_file(self, owner, hotel, guest):
        book(hotel, guest, date(2026, 11, 3))
        cancelled = book(hotel, guest, date(2026, 11, 9), status='cancelled', payment_status='pending')
        rows = rows_of(client_of(owner).get(URL, {'export': 'csv', 'status': 'cancelled'}))
        assert [row[0] for row in rows[1:]] == [cancelled.confirmation_code]
        rows = rows_of(client_of(owner).get(URL, {'export': 'csv', 'payment_status': 'pending'}))
        assert [row[0] for row in rows[1:]] == [cancelled.confirmation_code]

    def test_a_hotel_name_that_looks_like_a_formula_is_written_as_text(self, owner, hotel, guest):
        PropertyTranslation.objects.filter(property=hotel).update(name='=HYPERLINK("http://x")')
        book(hotel, guest, date(2026, 11, 3))
        rows = rows_of(client_of(owner).get(URL, {'export': 'csv'}))
        assert rows[1][BOOKINGS_EXPORT_HEADER.index('hotel')].startswith("'=")

    @override_settings(CSV_EXPORT_MAX_ROWS=1)
    def test_a_long_list_is_cut_and_says_so(self, owner, hotel, guest):
        book(hotel, guest, date(2026, 11, 3))
        book(hotel, guest, date(2026, 11, 9))
        rows = rows_of(client_of(owner).get(URL, {'export': 'csv'}))
        assert len(rows) == 3 and rows[-1] == ['truncated']

    def test_the_download_is_written_to_the_access_log(self, owner, hotel, guest):
        book(hotel, guest, date(2026, 11, 3))
        b''.join(client_of(owner).get(URL, {'export': 'csv'}).streaming_content)
        row = AdminAccessLog.objects.get(action='export_csv')
        assert (row.actor_id, row.details) == (owner.id, {'export': 'partner_bookings', 'rows': 1})


class TestTheExcelFile:
    def test_an_excel_file_with_the_same_columns(self, owner, hotel, guest):
        booking = book(hotel, guest, date(2026, 11, 3), rooms=2, price=Decimal('100000'))
        response = client_of(owner).get(URL, {'export': 'xlsx'})
        assert response.status_code == 200
        assert response['Content-Disposition'] == 'attachment; filename="partner_bookings.xlsx"'
        sheet = load_workbook(io.BytesIO(response.content)).active
        assert [cell.value for cell in sheet[1]] == BOOKINGS_EXPORT_HEADER
        assert sheet.cell(row=2, column=1).value == booking.confirmation_code
        assert sheet.cell(row=2, column=BOOKINGS_EXPORT_HEADER.index('total') + 1).value == 400000


class TestNothingElseChanges:
    def test_without_export_the_list_is_the_same_json_as_before(self, owner, hotel, guest):
        book(hotel, guest, date(2026, 11, 3))
        response = client_of(owner).get(URL)
        assert response.status_code == 200 and isinstance(response.data, list) and len(response.data) == 1

    def test_an_unknown_export_value_gives_the_json_list(self, owner, hotel, guest):
        book(hotel, guest, date(2026, 11, 3))
        assert isinstance(client_of(owner).get(URL, {'export': 'pdf'}).data, list)

    def test_a_guest_and_a_visitor_are_refused(self):
        assert client_of(make_guest()).get(URL, {'export': 'csv'}).status_code == 403
        assert APIClient().get(URL, {'export': 'csv'}).status_code in (401, 403)

    def test_staff_get_a_file_with_only_their_own_hotels(self, hotel, guest):
        book(hotel, guest, date(2026, 11, 3))
        rows = rows_of(client_of(make_staff()).get(URL, {'export': 'csv'}))
        assert rows == [BOOKINGS_EXPORT_HEADER]
